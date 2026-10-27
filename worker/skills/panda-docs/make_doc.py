#!/usr/bin/env python3
"""Markdown -> PDF / DOCX / PPTX with a clean, modern look.

usage: make_doc.py SOURCE.md OUT.(pdf|docx|pptx) [--theme ember|ocean|forest|mono] [--title T]
"""
from __future__ import annotations

import argparse
import io
import os
import re
import sys
import urllib.request

THEMES = {
    "ember": ("#E8643A", "#2A2622", "#FBE7DC"),
    "ocean": ("#2563EB", "#1F2937", "#DBEAFE"),
    "forest": ("#2F855A", "#1F2A24", "#DCF2E3"),
    "mono": ("#111111", "#222222", "#EEEEEE"),
}


# ── tiny Markdown block parser ─────────────────────────────────────────────
def parse(md: str):
    blocks, para, lines = [], [], md.replace("\r\n", "\n").split("\n")

    def flush():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush()
        elif m := re.match(r"^(#{1,3})\s+(.*)", s):
            flush()
            blocks.append((f"h{len(m.group(1))}", m.group(2)))
        elif re.match(r"^(-{3,}|\*{3,})$", s):
            flush()
            blocks.append(("hr", ""))
        elif m := re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s):
            flush()
            blocks.append(("img", (m.group(1), m.group(2))))
        elif s.lower().startswith("notes:"):
            flush()
            blocks.append(("notes", s[6:].strip()))
        elif s.startswith(">"):
            flush()
            blocks.append(("quote", s.lstrip("> ").strip()))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^([-*+]|\d+[.)])\s+", s):
            flush()
            items, ordered = [], bool(re.match(r"^\d", s))
            while i < len(lines) and re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]):
                items.append(re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", lines[i]))
                i += 1
            blocks.append(("ol" if ordered else "ul", items))
            continue
        else:
            para.append(ln)
        i += 1
    flush()
    return blocks


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def runs(text: str):
    """Split inline text into (text, bold, italic, code) runs."""
    out = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False, False))
        elif part.startswith("`"):
            out.append((part[1:-1], False, False, True))
        elif part.startswith("*"):
            out.append((part[1:-1], False, True, False))
        else:
            out.append((part, False, False, False))
    return out


def load_image(src: str) -> bytes | None:
    try:
        if src.startswith(("http://", "https://")):
            req = urllib.request.Request(src, headers={"User-Agent": "Mozilla/5.0 PandaDocs"})
            return urllib.request.urlopen(req, timeout=20).read()
        return open(os.path.expanduser(src), "rb").read()
    except Exception:
        return None


# ── PDF ────────────────────────────────────────────────────────────────────
# reportlab's built-in Helvetica has no Cyrillic, Greek or most accents: Ukrainian text came out
# as black boxes (■■■■). Embed a full Unicode TTF family instead; Helvetica only if none exists.
_FONT_DIRS = ["/usr/share/fonts/truetype", "/usr/share/fonts", "/usr/local/share/fonts",
              os.path.expanduser("~/.local/share/fonts"), os.path.expanduser("~/.fonts")]
_FAMILIES = [  # (regular, bold, italic, bold italic) file names, best first
    ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf"),
    ("NotoSans-Regular.ttf", "NotoSans-Bold.ttf", "NotoSans-Italic.ttf", "NotoSans-BoldItalic.ttf"),
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
    ("FreeSans.ttf", "FreeSansBold.ttf", "FreeSansOblique.ttf", "FreeSansBoldOblique.ttf"),
]
_MONOS = ["LiberationMono-Regular.ttf", "NotoSansMono-Regular.ttf", "DejaVuSansMono.ttf", "FreeMono.ttf"]


def _find_font(name):
    dirs = list(_FONT_DIRS)
    try:  # matplotlib ships DejaVu, so this works even on a bare machine with matplotlib installed
        import matplotlib
        dirs.append(os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf"))
    except Exception:
        pass
    for d in dirs:
        for root, _, files in os.walk(d):
            if name in files:
                return os.path.join(root, name)
    return None


def pdf_fonts():
    """Register a Unicode font family; returns (regular, bold, italic, mono) font names."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping
    names = ("PandaSans", "PandaSans-Bold", "PandaSans-Italic", "PandaSans-BoldItalic")
    for fam in _FAMILIES:
        paths = [_find_font(f) for f in fam]
        if not paths[0]:
            continue
        paths = [p or paths[0] for p in paths]  # a missing style falls back to regular
        try:
            for n, p in zip(names, paths):
                pdfmetrics.registerFont(TTFont(n, p))
            for bold, it, n in ((0, 0, names[0]), (1, 0, names[1]), (0, 1, names[2]), (1, 1, names[3])):
                addMapping("PandaSans", bold, it, n)
            mono = next((m for m in map(_find_font, _MONOS) if m), None)
            if mono:
                pdfmetrics.registerFont(TTFont("PandaMono", mono))
            return names[0], names[1], names[2], "PandaMono" if mono else names[0]
        except Exception:
            continue
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier"


def to_pdf(blocks, out, theme, title):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    REG, BOLD, ITAL, MONO = pdf_fonts()
    accent, ink, soft = (colors.HexColor(c) for c in THEMES[theme])
    base = ParagraphStyle("b", fontName=REG, fontSize=10.5, leading=15.5, textColor=ink, alignment=TA_LEFT)
    st = {
        "h1": ParagraphStyle("h1", parent=base, fontName=BOLD, fontSize=22, leading=27, spaceBefore=6, spaceAfter=10, textColor=accent),
        "h2": ParagraphStyle("h2", parent=base, fontName=BOLD, fontSize=15, leading=20, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", parent=base, fontName=BOLD, fontSize=12, leading=16, spaceBefore=8, spaceAfter=4),
        "p": ParagraphStyle("p", parent=base, spaceAfter=7),
        "quote": ParagraphStyle("q", parent=base, leftIndent=10, textColor=colors.HexColor("#6B6259"), fontName=ITAL, borderPadding=(4, 0, 4, 8)),
        "cell": ParagraphStyle("c", parent=base, fontSize=9.5, leading=13),
    }

    def rich(t):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
        t = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", t)
        t = re.sub(r"`([^`]+)`", rf"<font face='{MONO}'>\1</font>", t)
        return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"<link href='\2' color='#2563EB'>\1</link>", t)

    flow = []
    if title:
        flow += [Paragraph(rich(title), st["h1"]), Spacer(1, 4)]
    for kind, val in blocks:
        if kind in ("h1", "h2", "h3", "p", "quote"):
            flow.append(Paragraph(rich(val), st[kind]))
        elif kind in ("ul", "ol"):
            flow.append(ListFlowable([ListItem(Paragraph(rich(x), st["p"])) for x in val],
                                     bulletType="1" if kind == "ol" else "bullet", start="1" if kind == "ol" else None,
                                     bulletColor=accent, leftIndent=14))
        elif kind == "table" and val:
            data = [[Paragraph(rich(c), st["cell"]) for c in r] for r in val]
            w = (A4[0] - 40 * mm) / max(len(r) for r in val)
            tb = Table(data, colWidths=[w] * max(len(r) for r in val), repeatRows=1)
            tb.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), soft), ("FONTNAME", (0, 0), (-1, 0), BOLD),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D2C7")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            flow += [tb, Spacer(1, 8)]
        elif kind == "hr":
            flow.append(PageBreak())
        elif kind == "img":
            data = load_image(val[1])
            if data:
                img = Image(io.BytesIO(data))
                maxw = A4[0] - 40 * mm
                ratio = min(1.0, maxw / img.imageWidth, (110 * mm) / img.imageHeight)
                img.drawWidth, img.drawHeight = img.imageWidth * ratio, img.imageHeight * ratio
                flow += [img, Paragraph(f"<i>{rich(val[0])}</i>", st["cell"]) if val[0] else Spacer(1, 2), Spacer(1, 8)]

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(accent)
        canvas.rect(0, A4[1] - 6, A4[0], 6, stroke=0, fill=1)
        canvas.setFont(REG, 8)
        canvas.setFillColor(colors.HexColor("#9A9087"))
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(out, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm,
                      bottomMargin=20 * mm, title=title or os.path.basename(out)).build(flow, onFirstPage=deco, onLaterPages=deco)


# ── DOCX ───────────────────────────────────────────────────────────────────
def to_docx(blocks, out, theme, title):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches, Pt, RGBColor

    accent = RGBColor.from_string(THEMES[theme][0][1:])
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for lvl in (1, 2, 3):
        doc.styles[f"Heading {lvl}"].font.color.rgb = accent if lvl == 1 else RGBColor(0x2A, 0x26, 0x22)

    def add_runs(p, text):
        for t, b, it, code in runs(text):
            r = p.add_run(t)
            r.bold, r.italic = b, it
            if code:
                r.font.name = "Consolas"

    if title:
        doc.add_heading(title, 0)
    for kind, val in blocks:
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(val, int(kind[1]))
        elif kind == "p":
            add_runs(doc.add_paragraph(), val)
        elif kind == "quote":
            add_runs(doc.add_paragraph(style="Intense Quote"), val)
        elif kind in ("ul", "ol"):
            for x in val:
                add_runs(doc.add_paragraph(style="List Number" if kind == "ol" else "List Bullet"), x)
        elif kind == "table" and val:
            cols = max(len(r) for r in val)
            tb = doc.add_table(rows=len(val), cols=cols)
            tb.style = "Light Grid Accent 1"
            for ri, r in enumerate(val):
                for ci, c in enumerate(r):
                    cell = tb.cell(ri, ci)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], c)
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif kind == "hr":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "img":
            data = load_image(val[1])
            if data:
                doc.add_picture(io.BytesIO(data), width=Inches(6))
                if val[0]:
                    doc.add_paragraph(val[0]).runs[0].italic = True
    doc.save(out)


# ── PPTX ───────────────────────────────────────────────────────────────────
def to_pptx(blocks, out, theme, title):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    accent_hex, ink_hex, soft_hex = THEMES[theme]
    accent, ink, soft = (RGBColor.from_string(h[1:]) for h in (accent_hex, ink_hex, soft_hex))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    slides, cur = [], None
    for kind, val in blocks:
        if kind in ("h1", "h2"):
            cur = {"title": val, "bullets": [], "body": [], "img": None, "notes": "", "first": kind == "h1" and not slides}
            slides.append(cur)
        elif cur is None:
            cur = {"title": title or "", "bullets": [], "body": [], "img": None, "notes": "", "first": True}
            slides.append(cur)
        if cur is None or kind in ("h1", "h2"):
            continue
        if kind in ("ul", "ol"):
            cur["bullets"] += val
        elif kind in ("p", "quote", "h3"):
            cur["body"].append(val)
        elif kind == "img":
            cur["img"] = val
        elif kind == "notes":
            cur["notes"] += val + "\n"
        elif kind == "table" and val:
            cur["bullets"] += [" — ".join(r) for r in val[1:]] or [" — ".join(val[0])]

    def text_box(slide, x, y, w, h, text, size, color, bold=False):
        tf = slide.shapes.add_textbox(x, y, w, h).text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        for t, b, it, _ in runs(text):
            r = p.add_run()
            r.text, r.font.size, r.font.bold, r.font.italic = t, Pt(size), bold or b, it
            r.font.color.rgb = color
        return tf

    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if s["first"] or i == 0:
            band = slide.shapes.add_shape(1, 0, 0, W, H)
            band.fill.solid(); band.fill.fore_color.rgb = soft; band.line.fill.background()
            bar = slide.shapes.add_shape(1, Inches(0.8), Inches(3.05), Inches(1.2), Inches(0.12))
            bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
            text_box(slide, Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.6), s["title"], 44, ink, True)
            if s["body"]:
                text_box(slide, Inches(0.8), Inches(3.4), Inches(11), Inches(2), " ".join(s["body"]), 22, ink)
        else:
            bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), H)
            bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
            text_box(slide, Inches(0.7), Inches(0.45), Inches(12), Inches(1), s["title"], 32, ink, True)
            has_img = bool(s["img"] and load_image(s["img"][1]))
            tw = Inches(6.6) if has_img else Inches(11.8)
            y = Inches(1.6)
            if s["body"]:
                text_box(slide, Inches(0.7), y, tw, Inches(1.2), " ".join(s["body"]), 18, RGBColor(0x5A, 0x52, 0x4A))
                y = Inches(2.7)
            if s["bullets"]:
                tf = slide.shapes.add_textbox(Inches(0.7), y, tw, H - y - Inches(0.6)).text_frame
                tf.word_wrap = True
                for bi, b in enumerate(s["bullets"][:8]):
                    p = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
                    p.space_after = Pt(10)
                    dot = p.add_run(); dot.text = "●  "; dot.font.size = Pt(14); dot.font.color.rgb = accent
                    for t, bb, it, _ in runs(b):
                        r = p.add_run(); r.text, r.font.size, r.font.bold, r.font.italic = t, Pt(20), bb, it
                        r.font.color.rgb = ink
            if has_img:
                data = load_image(s["img"][1])
                pic = slide.shapes.add_picture(io.BytesIO(data), Inches(7.6), Inches(1.6), width=Inches(5.2))
                if pic.height > Inches(5.3):
                    ratio = Inches(5.3) / pic.height
                    pic.height, pic.width = Inches(5.3), Emu(int(pic.width * ratio))
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"].strip()
    prs.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("out")
    ap.add_argument("--theme", default="ember", choices=sorted(THEMES))
    ap.add_argument("--title", default="")
    a = ap.parse_args()
    md = open(os.path.expanduser(a.source), encoding="utf-8").read() if a.source != "-" else sys.stdin.read()
    out = os.path.expanduser(a.out)
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    blocks = parse(md)
    ext = out.rsplit(".", 1)[-1].lower()
    {"pdf": to_pdf, "docx": to_docx, "pptx": to_pptx}[ext](blocks, out, a.theme, a.title)
    print(f"saved {out} ({os.path.getsize(out) // 1024} KiB)")


if __name__ == "__main__":
    main()
